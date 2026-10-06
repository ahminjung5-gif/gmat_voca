"""
GMAT_VOCA_RC_유의어_리스트.xlsx 의 'Day별 학습' 시트를 data/rcsets.js 로 변환합니다.

사용법:  pip install openpyxl
         python scripts/convert_rc.py GMAT_VOCA_RC_유의어_리스트.xlsx
"""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

SHEET = "Day별 학습"
src = Path(sys.argv[1] if len(sys.argv) > 1 else "GMAT_VOCA_RC_유의어_리스트.xlsx")
out = Path(__file__).resolve().parent.parent / "data" / "rcsets.js"

def clean(v):
    return "" if v is None else str(v).replace("\\n", "\n").strip()

def senses(text):
    # "(1) ...\n(2) ..." -> ["...", "..."]
    lines = [l.strip() for l in clean(text).split("\n") if l.strip()]
    out, cur = [], None
    for l in lines:
        m = re.match(r"^\((\d+)\)\s*(.*)$", l)
        if m:
            out.append(m.group(2).strip())
        elif out:
            out[-1] += " " + l          # 번호 없이 이어지는 줄은 앞 뜻에 붙임
        else:
            out.append(l)
    return out

wb = load_workbook(src, read_only=True)
rows = list(wb[SHEET].iter_rows(values_only=True))
header = [clean(h) for h in rows[0]]
ix = {n: header.index(n) for n in ["DAY", "단어", "난이도", "한글 뜻", "동의어", "등장 횟수", "메모/예문"]}

items = []
for r in rows[1:]:
    w = clean(r[ix["단어"]])
    m = re.search(r"\d+", clean(r[ix["DAY"]]))
    if not w or not m:
        continue
    means, syns = senses(r[ix["한글 뜻"]]), senses(r[ix["동의어"]])
    n = max(len(means), len(syns))
    means += [""] * (n - len(means)); syns += [""] * (n - len(syns))
    items.append({
        "d": int(m.group()),
        "w": w,
        "lv": clean(r[ix["난이도"]]),
        "c": int(r[ix["등장 횟수"]] or 0),
        "s": [[a, b] for a, b in zip(means, syns)],   # [뜻, 유의어] 묶음
        "e": clean(r[ix["메모/예문"]]),
    })

out.write_text("window.RCSETS=" + json.dumps(items, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(f"{len(items)} sets -> {out}")
