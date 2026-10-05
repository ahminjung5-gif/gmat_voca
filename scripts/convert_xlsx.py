"""
GMAT_VOCA.xlsx 의 'Day별 학습' 시트를 data/words.js 로 변환합니다.
엑셀을 수정한 뒤 이 스크립트를 다시 실행하면 게임 데이터가 갱신됩니다.

사용법:  pip install openpyxl
         python scripts/convert_xlsx.py GMAT_VOCA.xlsx
"""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook

SHEET = "Day별 학습"
src = Path(sys.argv[1] if len(sys.argv) > 1 else "GMAT_VOCA.xlsx")
out = Path(__file__).resolve().parent.parent / "data" / "words.js"

def clean(v):
    if v is None:
        return ""
    return str(v).strip()

def clean_pron(p):
    p = clean(p)
    # 엑셀 원본의 깨진 IPA 문자를 표준 기호로 교체
    return p.replace("ӕ", "æ").replace("Ʒ", "ʒ").replace("|", "ˈ")

# 원본 DERIVATIVE 열에 두 단어가 붙어 있는 경우가 있음 (예: "suspiciousskepticism")
# 표제어 앞 4글자가 단어 중간에서 시작하면 그 지점에서 ", " 로 분리
PREFIXES = {"un", "re", "dis", "over", "under", "non", "inter", "counter", "mis", "pre", "in", "im",
            "il", "ir", "de", "co", "sub", "super", "out", "self", "anti", "multi", "semi", "trans",
            "post", "fore", "up", "down", "well", "ill", "en", "em", "ex", "mid", "pro", "con", "com", "be"}

def fix_glue(word, deriv):
    w = word.lower()
    if not deriv or len(w) < 4 or not re.fullmatch(r"[a-z]+", w):
        return deriv
    stem = w[:4]

    def split(m):
        tok = m.group(0)
        i = tok.lower().find(stem, 1)
        if i < 3:
            return tok
        left, right = tok[:i], tok[i:]
        # "carefulun" + "scrupulous" -> "careful" + "unscrupulous"
        for pre in sorted(PREFIXES, key=len, reverse=True):
            if left.lower().endswith(pre) and len(left) - len(pre) >= 3:
                left, right = left[:-len(pre)], left[-len(pre):] + right
                break
        if left.lower() in PREFIXES:
            return tok
        return f"{left}, {right}"

    return re.sub(r"[A-Za-z]+", split, deriv)

wb = load_workbook(src, read_only=True)
ws = wb[SHEET]
rows = list(ws.iter_rows(values_only=True))
header = [clean(h) for h in rows[0]]
idx = {n: header.index(n) for n in ["DAY", "구분", "WORD", "발음", "MEANING", "DERIVATIVE", "EXAMPLE"]}

words = []
for r in rows[1:]:
    w = clean(r[idx["WORD"]])
    m = re.search(r"\d+", clean(r[idx["DAY"]]))
    if not w or not m:
        continue
    words.append({
        "d": int(m.group()),              # day
        "s": clean(r[idx["구분"]]),        # RC / CR / 기본
        "w": w,                           # word
        "p": clean_pron(r[idx["발음"]]),    # pronunciation
        "m": clean(r[idx["MEANING"]]),     # meaning
        "v": fix_glue(w, clean(r[idx["DERIVATIVE"]])),  # derivative
        "e": clean(r[idx["EXAMPLE"]]),     # example
    })

out.parent.mkdir(parents=True, exist_ok=True)
out.write_text("window.WORDS=" + json.dumps(words, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(f"{len(words)} words -> {out}")
